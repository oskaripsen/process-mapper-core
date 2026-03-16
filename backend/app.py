from contextlib import asynccontextmanager
from datetime import datetime
from typing import Any, Dict, List, Optional, Set, Tuple
import json
import logging
import os
import uuid

from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr
import uvicorn

from db import Database
from schemas.process_flow import ProcessFlow, ProcessNode, ProcessEdge, NodeType, AutomationLevel, FlowPatch
from services.chat_service import ChatService
from services.whisper_service import WhisperService
from utils.auth import create_access_token, get_current_user_id, hash_password, verify_password, verify_token

load_dotenv()
load_dotenv(os.path.join(os.path.dirname(__file__), "..", ".env"))

logger = logging.getLogger(__name__)

from utils.storage import StorageService

db = Database()
whisper_service = WhisperService()
storage_service = StorageService()

# Lazy-initialised services (require OPENAI_API_KEY at runtime, not import time)
_chat_service = None
_llm_service = None
_intent_translator = None
_patch_engine = None
_reactflow_translator = None
_sop_ai_service = None


def get_chat_service():
    global _chat_service
    if _chat_service is None:
        _chat_service = ChatService()
    return _chat_service


def get_llm_service():
    global _llm_service
    if _llm_service is None:
        from services.llm_service import LLMService
        _llm_service = LLMService()
    return _llm_service


def get_intent_translator():
    global _intent_translator
    if _intent_translator is None:
        from services.intent_translator import IntentTranslator
        _intent_translator = IntentTranslator()
    return _intent_translator


def get_patch_engine():
    global _patch_engine
    if _patch_engine is None:
        from services.patch_engine import PatchEngine
        _patch_engine = PatchEngine()
    return _patch_engine


def get_reactflow_translator():
    global _reactflow_translator
    if _reactflow_translator is None:
        from translators.reactflow_translator import ReactFlowTranslator
        _reactflow_translator = ReactFlowTranslator()
    return _reactflow_translator


def get_sop_ai_service():
    global _sop_ai_service
    if _sop_ai_service is None:
        from services.sop_ai_service import SOPAIService
        _sop_ai_service = SOPAIService()
    return _sop_ai_service


def reactflow_to_processflow(reactflow_data: dict) -> ProcessFlow:
    """Convert React Flow format to ProcessFlow canonical schema."""
    nodes = []
    for rf_node in reactflow_data.get("nodes", []):
        node_data = rf_node.get("data", rf_node)

        node_type = NodeType.PROCESS
        rf_type = rf_node.get("type", node_data.get("type", "default"))
        if rf_type == "start":
            node_type = NodeType.START
        elif rf_type == "end":
            node_type = NodeType.END
        elif rf_type == "decision":
            node_type = NodeType.DECISION
        elif rf_type == "merge":
            node_type = NodeType.MERGE

        automation_value = node_data.get("manualOrAutomated") or rf_node.get("manualOrAutomated") or "manual"
        automation_level = (
            AutomationLevel.AUTOMATED if str(automation_value).lower().strip() == "automated"
            else AutomationLevel.MANUAL
        )

        nodes.append(ProcessNode(
            id=rf_node["id"],
            type=node_type,
            label=node_data.get("label", rf_node.get("label", "Unknown")),
            owner=node_data.get("owner", rf_node.get("owner")),
            system=node_data.get("system", rf_node.get("system")),
            automation=automation_level,
            logical_id=node_data.get("logical_id", rf_node.get("logical_id")),
            user_modified=node_data.get("user_modified", rf_node.get("user_modified", False)),
            screenshot_url=node_data.get("screenshot_url", rf_node.get("screenshot_url")),
        ))

    node_ids = {n.id for n in nodes}
    edges = []
    for rf_edge in reactflow_data.get("edges", []):
        source_id = rf_edge.get("source")
        target_id = rf_edge.get("target")
        if source_id not in node_ids or target_id not in node_ids:
            continue
        edges.append(ProcessEdge(
            id=rf_edge.get("id", f"edge-{source_id}-{target_id}"),
            source=source_id,
            target=target_id,
            condition=rf_edge.get("condition") or rf_edge.get("label"),
        ))

    return ProcessFlow(nodes=nodes, edges=edges)


class RegisterRequest(BaseModel):
    email: EmailStr
    username: str
    password: str


class LoginRequest(BaseModel):
    identifier: Optional[str] = None
    email: Optional[EmailStr] = None
    password: str


class TaxonomyCreate(BaseModel):
    name: str
    description: Optional[str] = ""
    code: Optional[str] = ""
    level: int = 0
    parent_id: Optional[str] = None
    sort_order: Optional[int] = None


class TaxonomyUpdate(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    code: Optional[str] = None
    sort_order: Optional[int] = None


class TaxonomyMoveRequest(BaseModel):
    parent_id: Optional[str] = None
    sort_order: Optional[int] = None


class TaxonomyReorderItem(BaseModel):
    id: str
    sort_order: int


class TaxonomyReorderRequest(BaseModel):
    items: List[TaxonomyReorderItem]


class ProcessAssignmentCreateRequest(BaseModel):
    process_id: str
    user_email: EmailStr
    role: str


class FlowCreate(BaseModel):
    process_id: str
    title: str
    description: Optional[str] = ""
    flow_data: Dict[str, Any]


class FlowUpdate(BaseModel):
    title: Optional[str] = None
    description: Optional[str] = None
    flow_data: Dict[str, Any]


class IncrementalFlowRequest(BaseModel):
    transcript: str
    accumulatedTranscript: Optional[str] = None
    existingFlow: Optional[Dict[str, Any]] = None
    sessionId: Optional[Any] = None
    validationErrors: Optional[List[str]] = None
    userEdits: Optional[List[Dict[str, Any]]] = None
    extractedIntent: Optional[Dict[str, Any]] = None


VALID_ASSIGNMENT_ROLES: Set[str] = {"owner", "delegator", "delegatee"}


def _normalize_role(role: str) -> str:
    normalized = (role or "").strip().lower()
    if normalized not in VALID_ASSIGNMENT_ROLES:
        raise HTTPException(status_code=400, detail="Invalid role. Must be owner, delegator, or delegatee")
    return normalized


async def _get_user_email(conn, user_id: str) -> Optional[str]:
    async with conn.execute("SELECT email FROM users WHERE id = ?", (user_id,)) as cursor:
        row = await cursor.fetchone()
    return row["email"] if row else None


async def _get_process_row(conn, process_id: str) -> Optional[Dict[str, Any]]:
    async with conn.execute(
        "SELECT id, user_id, parent_id, level, name FROM process_taxonomy WHERE id = ?",
        (process_id,),
    ) as cursor:
        row = await cursor.fetchone()
    return dict(row) if row else None


async def _get_access_request_email(conn, process_id: Optional[str]) -> Optional[str]:
    if not process_id:
        return None

    async with conn.execute(
        """
        SELECT pa.user_email
        FROM process_assignments pa
        WHERE pa.process_id = ? AND pa.role = 'owner' AND pa.is_active = 1
        ORDER BY pa.assigned_at ASC
        LIMIT 1
        """,
        (process_id,),
    ) as cursor:
        row = await cursor.fetchone()
    if row and row["user_email"]:
        return row["user_email"]

    process_row = await _get_process_row(conn, process_id)
    if not process_row:
        return None
    return await _get_user_email(conn, process_row["user_id"])


async def _raise_level_access_denied(
    conn,
    reference_process_id: Optional[str],
    action: str,
) -> None:
    reference = await _get_process_row(conn, reference_process_id) if reference_process_id else None
    owner_email = await _get_access_request_email(conn, reference_process_id)
    detail = (
        f"This is not allowed. {action} can only be done if you have owner or delegator rights "
        f"one level above."
    )
    if reference and reference.get("level") is not None:
        detail += f" Required access: L{reference['level']}."
    if owner_email:
        detail += f" Request access from {owner_email}."
    raise HTTPException(status_code=403, detail=detail)


async def _ensure_creator_owner_assignment(conn, process_id: str) -> None:
    process_row = await _get_process_row(conn, process_id)
    if not process_row:
        return

    creator_user_id = process_row["user_id"]
    creator_email = await _get_user_email(conn, creator_user_id)
    if not creator_email:
        return

    async with conn.execute(
        """
        SELECT id FROM process_assignments
        WHERE process_id = ? AND role = 'owner' AND is_active = 1
          AND (user_id = ? OR lower(user_email) = lower(?))
        LIMIT 1
        """,
        (process_id, creator_user_id, creator_email),
    ) as cursor:
        existing = await cursor.fetchone()
    if existing:
        return

    assignment_id = str(uuid.uuid4())
    await conn.execute(
        """
        INSERT INTO process_assignments (id, process_id, user_email, user_id, role, assigned_by, is_active)
        VALUES (?, ?, ?, ?, 'owner', ?, 1)
        """,
        (assignment_id, process_id, creator_email.lower(), creator_user_id, creator_user_id),
    )


async def _get_effective_assignments(conn, process_id: str) -> List[Dict[str, Any]]:
    query = """
        WITH RECURSIVE lineage(id, parent_id, depth) AS (
            SELECT id, parent_id, 0
            FROM process_taxonomy
            WHERE id = ?
            UNION ALL
            SELECT pt.id, pt.parent_id, lineage.depth + 1
            FROM process_taxonomy pt
            JOIN lineage ON pt.id = lineage.parent_id
        )
        SELECT
            pa.id,
            pa.process_id,
            pa.user_email,
            pa.user_id,
            pa.role,
            pa.assigned_by,
            pa.assigned_at,
            pa.is_active,
            pt.name AS process_name,
            lineage.depth
        FROM lineage
        JOIN process_assignments pa ON pa.process_id = lineage.id AND pa.is_active = 1
        LEFT JOIN process_taxonomy pt ON pt.id = pa.process_id
        ORDER BY lineage.depth ASC, pa.assigned_at ASC
    """
    async with conn.execute(query, (process_id,)) as cursor:
        rows = await cursor.fetchall()
    assignments: List[Dict[str, Any]] = []
    seen_identities: Set[str] = set()
    for row in rows:
        item = dict(row)
        item.pop("depth", None)
        identity = (
            f"id:{item['user_id']}"
            if item.get("user_id")
            else f"email:{(item.get('user_email') or '').lower()}"
        )
        # Keep closest assignment for each user (direct before inherited),
        # which prevents duplicated entries when creator is owner at multiple levels.
        if identity in seen_identities:
            continue
        seen_identities.add(identity)
        assignments.append(item)
    return assignments


async def _get_requester_role_context(conn, process_id: str, user_id: str) -> Dict[str, Any]:
    process_row = await _get_process_row(conn, process_id)
    if not process_row:
        raise HTTPException(status_code=404, detail="Process not found")

    await _ensure_creator_owner_assignment(conn, process_id)
    requester_email = await _get_user_email(conn, user_id)
    assignments = await _get_effective_assignments(conn, process_id)
    requester_assignments = [
        a for a in assignments
        if a.get("user_id") == user_id
        or (
            requester_email
            and a.get("user_email")
            and a["user_email"].lower() == requester_email.lower()
        )
    ]

    is_creator = process_row["user_id"] == user_id
    requester_roles = {a.get("role") for a in requester_assignments}
    is_owner = is_creator or "owner" in requester_roles
    is_delegator = "delegator" in requester_roles
    has_access = is_creator or len(requester_assignments) > 0
    can_manage = is_owner or is_delegator

    return {
        "process": process_row,
        "assignments": assignments,
        "has_access": has_access,
        "is_owner": is_owner,
        "is_delegator": is_delegator,
        "can_manage": can_manage,
    }


async def _can_view_process_in_tree(conn, process_id: str, user_id: str) -> bool:
    user_email = await _get_user_email(conn, user_id)
    async with conn.execute(
        """
        WITH RECURSIVE
        seed(id) AS (
            SELECT id
            FROM process_taxonomy
            WHERE user_id = ?
            UNION
            SELECT process_id
            FROM process_assignments
            WHERE is_active = 1
              AND (
                  user_id = ?
                  OR (? IS NOT NULL AND lower(user_email) = lower(?))
              )
        ),
        down(id) AS (
            SELECT id FROM seed
            UNION
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN down d ON pt.parent_id = d.id
        ),
        up(id, parent_id) AS (
            SELECT pt.id, pt.parent_id
            FROM process_taxonomy pt
            WHERE pt.id IN (SELECT id FROM down)
            UNION
            SELECT pt.id, pt.parent_id
            FROM process_taxonomy pt
            JOIN up ON up.parent_id = pt.id
        )
        SELECT 1
        FROM up
        WHERE id = ?
        LIMIT 1
        """,
        (user_id, user_id, user_email, user_email, process_id),
    ) as cursor:
        row = await cursor.fetchone()
    return row is not None


async def _next_sibling_sort_order(
    conn,
    user_id: str,
    parent_id: Optional[str],
    exclude_id: Optional[str] = None,
) -> int:
    base_query = "SELECT MAX(sort_order) AS max_sort_order FROM process_taxonomy WHERE user_id = ?"
    params: List[Any] = [user_id]
    if parent_id is None:
        base_query += " AND parent_id IS NULL"
    else:
        base_query += " AND parent_id = ?"
        params.append(parent_id)
    if exclude_id:
        base_query += " AND id != ?"
        params.append(exclude_id)
    async with conn.execute(base_query, tuple(params)) as cursor:
        row = await cursor.fetchone()
    max_sort_order = row["max_sort_order"] if row else None
    return (max_sort_order or 0) + 1


async def _subtree_height(conn, user_id: str, root_id: str) -> int:
    async with conn.execute(
        "SELECT id, parent_id FROM process_taxonomy WHERE user_id = ?",
        (user_id,),
    ) as cursor:
        rows = await cursor.fetchall()
    children_by_parent: Dict[Optional[str], List[str]] = {}
    for row in rows:
        children_by_parent.setdefault(row["parent_id"], []).append(row["id"])

    def _dfs(node_id: str) -> int:
        children = children_by_parent.get(node_id, [])
        if not children:
            return 0
        return 1 + max(_dfs(child_id) for child_id in children)

    return _dfs(root_id)


async def _update_subtree_levels(conn, user_id: str, root_id: str, level: int) -> None:
    await conn.execute(
        "UPDATE process_taxonomy SET level = ? WHERE id = ? AND user_id = ?",
        (level, root_id, user_id),
    )
    async with conn.execute(
        "SELECT id FROM process_taxonomy WHERE user_id = ? AND parent_id = ?",
        (user_id, root_id),
    ) as cursor:
        child_rows = await cursor.fetchall()
    for child in child_rows:
        await _update_subtree_levels(conn, user_id, child["id"], level + 1)


async def init_schema():
    conn = db.get_connection()
    schema_path = os.path.join(os.path.dirname(__file__), "schema.sql")
    with open(schema_path, "r", encoding="utf-8") as f:
        await conn.executescript(f.read())

    # Backward-compatible additive migrations for existing DBs.
    sop_migrations = [
        ("process_taxonomy", "sort_order", "INTEGER DEFAULT 0"),
        ("process_taxonomy", "sop_status", "TEXT"),
        ("process_taxonomy", "sop_draft_url", "TEXT"),
        ("process_taxonomy", "sop_final_url", "TEXT"),
        ("process_taxonomy", "sop_final_pdf_url", "TEXT"),
        ("process_taxonomy", "sop_generated_at", "TEXT"),
        ("process_taxonomy", "sop_finalized_at", "TEXT"),
        ("process_flow_versions", "sop_url", "TEXT"),
        ("process_flow_versions", "version_type", "TEXT"),
    ]
    for table, col, col_type in sop_migrations:
        try:
            await conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {col_type}")
        except Exception:
            pass
    await conn.commit()


@asynccontextmanager
async def lifespan(_app: FastAPI):
    await db.initialize(os.getenv("DATABASE_PATH", "./data/process_mapper.db"))
    await init_schema()
    yield
    await db.close()


app = FastAPI(title="Process Mapper Core API", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


@app.get("/")
async def root():
    return {"message": "Process Mapper Core API"}


@app.get("/health")
async def health():
    return {"status": "ok", "time": datetime.utcnow().isoformat()}


@app.post("/api/auth/register")
async def register(payload: RegisterRequest):
    conn = db.get_connection()
    user_id = str(uuid.uuid4())
    try:
        await conn.execute(
            "INSERT INTO users (id, email, username, password_hash) VALUES (?, ?, ?, ?)",
            (user_id, payload.email.lower(), payload.username, hash_password(payload.password)),
        )
        await conn.commit()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Registration failed: {exc}") from exc

    token = create_access_token(user_id, payload.email.lower(), payload.username)
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {"id": user_id, "email": payload.email.lower(), "username": payload.username},
    }


@app.post("/api/auth/login")
async def login(payload: LoginRequest):
    login_identifier = (payload.identifier or payload.email or "").strip()
    if not login_identifier:
        raise HTTPException(status_code=400, detail="Email or username is required")

    normalized_identifier = login_identifier.lower()
    conn = db.get_connection()
    async with conn.execute(
        """
        SELECT id, email, username, password_hash FROM users
        WHERE lower(email) = ? OR lower(username) = ?
        LIMIT 1
        """,
        (normalized_identifier, normalized_identifier),
    ) as cursor:
        user = await cursor.fetchone()
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid credentials")

    token = create_access_token(user["id"], user["email"], user["username"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {"id": user["id"], "email": user["email"], "username": user["username"]},
    }


@app.get("/api/auth/me")
async def me(user: Dict[str, Any] = Depends(verify_token)):
    return user


@app.post("/transcribe")
async def transcribe_audio(file: UploadFile = File(...), user: Dict[str, Any] = Depends(verify_token)):
    content = await file.read()
    result = await whisper_service.transcribe_audio_from_bytes(content, file.filename or "audio.wav")
    return {"transcript": result.get("text", ""), "segments": result.get("segments", []), "words": result.get("words", [])}


@app.post("/generate-flow")
async def generate_flow(request: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    transcript = request.get("transcript", "")
    return {
        "nodes": [{"id": "start", "position": {"x": 100, "y": 100}, "data": {"label": "Start"}, "type": "start"}],
        "edges": [],
        "metadata": {"source": "core", "transcript_length": len(transcript)},
    }


@app.post("/generate-incremental-flow")
async def generate_incremental_flow(request: IncrementalFlowRequest, user: Dict[str, Any] = Depends(verify_token)):
    try:
        if not request.transcript.strip():
            raise HTTPException(status_code=400, detail="Transcript cannot be empty")

        existing_flow_data = request.existingFlow or {"nodes": [], "edges": []}

        # Apply user deletions before processing
        if request.userEdits:
            deleted_node_ids = {
                e.get("id") for e in request.userEdits if e.get("type") == "node_deleted"
            }
            deleted_edge_ids = {
                e.get("id") for e in request.userEdits if e.get("type") == "edge_deleted"
            }
            if deleted_node_ids:
                existing_flow_data["nodes"] = [
                    n for n in existing_flow_data.get("nodes", []) if n.get("id") not in deleted_node_ids
                ]
            if deleted_edge_ids:
                existing_flow_data["edges"] = [
                    e for e in existing_flow_data.get("edges", []) if e.get("id") not in deleted_edge_ids
                ]

        # Filter stale edges
        existing_node_ids = {n.get("id") for n in existing_flow_data.get("nodes", [])}
        existing_flow_data["edges"] = [
            e for e in existing_flow_data.get("edges", [])
            if e.get("source") in existing_node_ids and e.get("target") in existing_node_ids
        ]

        current_flow = reactflow_to_processflow(existing_flow_data)

        from utils.context_builder import build_flow_context
        context = build_flow_context(
            existing_flow=existing_flow_data,
            validation_errors=request.validationErrors,
            user_edits=request.userEdits,
            accumulated_transcript=request.accumulatedTranscript,
        )

        llm_service = get_llm_service()
        intent_translator = get_intent_translator()
        patch_engine = get_patch_engine()
        translator = get_reactflow_translator()

        # Extract intent from transcript
        intent = await llm_service.extract_process_intent(request.transcript, context)
        if not intent:
            if len(current_flow.nodes) > 0:
                existing_positions = {
                    n.get("id"): n.get("position")
                    for n in existing_flow_data.get("nodes", []) if n.get("position")
                }
                algo = "relative" if existing_positions else "hierarchical"
                return translator.translate(current_flow, layout_algorithm=algo, existing_positions=existing_positions, auto_layout=False)
            raise HTTPException(status_code=400, detail="No business process information found in transcript.")

        # Translate intent to patch
        try:
            patch = intent_translator.translate_intent_to_flow_patch(intent, existing_flow=current_flow)
        except Exception as exc:
            if "no process steps" in str(exc).lower():
                patch = FlowPatch(operations=[], source="intent_translator_empty")
            else:
                raise HTTPException(status_code=500, detail=f"Failed to translate intent: {exc}")

        # If no operations, return existing flow as-is
        if not patch or not patch.operations:
            if len(current_flow.nodes) > 0:
                existing_positions = {
                    n.get("id"): n.get("position")
                    for n in existing_flow_data.get("nodes", []) if n.get("position")
                }
                algo = "relative" if existing_positions else "hierarchical"
                return translator.translate(current_flow, layout_algorithm=algo, existing_positions=existing_positions, auto_layout=False)
            raise HTTPException(status_code=400, detail="No business process information found in transcript.")

        # Apply patch
        updated_flow, validation_errors = await run_in_threadpool(patch_engine.apply_patch, current_flow, patch)

        # Collect existing positions for layout
        existing_positions = {}
        for n in existing_flow_data.get("nodes", []):
            nid = n.get("id")
            if nid and n.get("position"):
                existing_positions[nid] = n["position"]

        new_node_ids = {n.id for n in updated_flow.nodes if n.id not in existing_positions}

        react_flow_data = translator.translate(
            updated_flow, layout_algorithm="manual", existing_positions=existing_positions, auto_layout=False,
        )

        # Apply layout
        from services.layout_service import layout_service
        layout_nodes = react_flow_data.get("nodes", [])
        layout_edges = react_flow_data.get("edges", [])

        has_existing_layout = len(existing_positions) > 0 and any(
            p.get("x", 0) != 0 or p.get("y", 0) != 0 for p in existing_positions.values()
        )

        if has_existing_layout and new_node_ids and len(new_node_ids) < len(layout_nodes):
            positioned_nodes, routed_edges = layout_service.apply_incremental_layout(
                layout_nodes, layout_edges, new_node_ids, existing_positions,
            )
        else:
            positioned_nodes, routed_edges = await run_in_threadpool(
                layout_service.apply_full_layout, layout_nodes, layout_edges, False,
            )

        react_flow_data["nodes"] = positioned_nodes
        react_flow_data["edges"] = routed_edges

        if validation_errors:
            react_flow_data["validationErrors"] = validation_errors

        return react_flow_data

    except HTTPException:
        raise
    except Exception as exc:
        logger.error(f"generate-incremental-flow error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/upload-doc")
async def upload_doc(
    files: List[UploadFile] = File(...),
    contexts: Optional[str] = Form(None),
    existing_flow: Optional[str] = Form(None),
    user: Dict[str, Any] = Depends(verify_token),
):
    try:
        if not files:
            raise HTTPException(status_code=400, detail="No files provided")

        file_contexts: Dict[str, str] = {}
        if contexts:
            try:
                file_contexts = json.loads(contexts)
            except json.JSONDecodeError:
                pass

        from utils.file_extraction import extract_text, extract_images

        all_text_parts: List[str] = []
        all_images: List[str] = []

        for f in files:
            file_context = file_contexts.get(f.filename or "", "")
            try:
                text = await extract_text(f)
                if text.strip():
                    header = f"--- From {f.filename} ---"
                    if file_context:
                        header += f"\nContext: {file_context}"
                    all_text_parts.append(f"{header}\n{text}")
            except Exception as exc:
                logger.warning(f"Text extraction failed for {f.filename}: {exc}")

            await f.seek(0)
            try:
                images = await extract_images(f)
                if images:
                    all_images.extend(images)
            except Exception as exc:
                logger.warning(f"Image extraction failed for {f.filename}: {exc}")

        combined_text = "\n\n".join(all_text_parts) if all_text_parts else ""
        if not combined_text.strip() and not all_images:
            raise HTTPException(status_code=400, detail="No content could be extracted from the uploaded files.")

        existing_flow_data = None
        if existing_flow:
            try:
                existing_flow_data = json.loads(existing_flow)
            except json.JSONDecodeError:
                pass

        from utils.context_builder import build_flow_context
        context = build_flow_context(existing_flow=existing_flow_data)

        llm_service = get_llm_service()
        intent_translator = get_intent_translator()
        patch_engine = get_patch_engine()
        translator = get_reactflow_translator()

        intent = await llm_service.extract_process_intent_from_documents(
            combined_text,
            context=context,
            images=all_images if all_images else None,
        )
        if not intent:
            raise HTTPException(status_code=400, detail="No business process found in the uploaded documents.")

        current_flow = reactflow_to_processflow(existing_flow_data) if existing_flow_data else ProcessFlow(nodes=[], edges=[])

        try:
            patch = intent_translator.translate_intent_to_flow_patch(intent, existing_flow=current_flow)
        except Exception as exc:
            if "no process steps" in str(exc).lower():
                raise HTTPException(status_code=400, detail="No process steps detected in the uploaded documents.")
            raise HTTPException(status_code=500, detail=f"Failed to translate intent: {exc}")

        if not patch or not patch.operations:
            raise HTTPException(status_code=400, detail="No business process found in the uploaded documents.")

        updated_flow, validation_errors = await run_in_threadpool(patch_engine.apply_patch, current_flow, patch)

        react_flow_data = translator.translate(updated_flow, layout_algorithm="manual", existing_positions=None, auto_layout=False)

        from services.layout_service import layout_service
        positioned_nodes, routed_edges = await run_in_threadpool(
            layout_service.apply_full_layout,
            react_flow_data.get("nodes", []),
            react_flow_data.get("edges", []),
            False,
        )
        react_flow_data["nodes"] = positioned_nodes
        react_flow_data["edges"] = routed_edges

        if validation_errors:
            react_flow_data["validationErrors"] = validation_errors

        return react_flow_data

    except HTTPException:
        raise
    except Exception as exc:
        logger.error(f"upload-doc error: {exc}")
        raise HTTPException(status_code=500, detail=str(exc))


@app.post("/auto-layout")
async def auto_layout(payload: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    existing_flow_data = payload.get("existingFlow", {})
    if not existing_flow_data or not existing_flow_data.get("nodes"):
        raise HTTPException(status_code=400, detail="Existing flow with nodes is required")

    user_edits = payload.get("userEdits")
    if user_edits:
        deleted_node_ids = {e.get("id") for e in user_edits if e.get("type") == "node_deleted"}
        deleted_edge_ids = {e.get("id") for e in user_edits if e.get("type") == "edge_deleted"}
        if deleted_node_ids:
            existing_flow_data["nodes"] = [
                n for n in existing_flow_data.get("nodes", []) if n.get("id") not in deleted_node_ids
            ]
        if deleted_edge_ids:
            existing_flow_data["edges"] = [
                e for e in existing_flow_data.get("edges", []) if e.get("id") not in deleted_edge_ids
            ]

    node_ids = {n.get("id") for n in existing_flow_data.get("nodes", [])}
    existing_flow_data["edges"] = [
        e for e in existing_flow_data.get("edges", [])
        if e.get("source") in node_ids and e.get("target") in node_ids
    ]

    current_flow = reactflow_to_processflow(existing_flow_data)
    translator = get_reactflow_translator()

    existing_positions = {}
    for n in existing_flow_data.get("nodes", []):
        nid = n.get("id")
        if nid and "position" in n:
            existing_positions[nid] = n["position"]

    react_flow_data = translator.translate(
        current_flow, layout_algorithm="manual", existing_positions=existing_positions, auto_layout=False,
    )

    from services.layout_service import layout_service
    positioned_nodes, routed_edges = await run_in_threadpool(
        layout_service.apply_full_layout,
        react_flow_data.get("nodes", []),
        react_flow_data.get("edges", []),
        False,
    )
    react_flow_data["nodes"] = positioned_nodes
    react_flow_data["edges"] = routed_edges
    return react_flow_data


@app.post("/api/chat")
async def chat(payload: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    return {"response": f"Core assistant received: {payload.get('message', '')}"}


@app.post("/api/clarify-process")
async def clarify_process(payload: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    return {"questions": [], "summary": payload.get("message", "")}


@app.post("/api/clarify-process/stream")
async def clarify_process_stream(payload: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    messages_raw = payload.get("messages", [])
    existing_flow = payload.get("existingFlow")

    is_first_flow = (
        not existing_flow
        or not existing_flow.get("nodes")
        or len(existing_flow.get("nodes", [])) == 0
    )

    if is_first_flow:
        system_content = (
            "You are a process analyst. The user is describing their business process for the FIRST TIME. "
            "You will help them clarify and understand the complete process.\n\n"
            "After understanding the process, format your response as follows:\n"
            "**Process Flow**: I've identified your process with [X] process steps, [Y] decision points (if any), and [Z] merge points (if any). "
            "[Brief summary of the overall process flow, including key actors and systems]\n\n"
            "**Clarifying Questions**: [If anything is unclear, ask 1-3 short clarifying questions as bullet points]\n\n"
            "Keep your response concise and focused on understanding the complete process. Use markdown formatting for emphasis."
        )
    else:
        system_content = (
            "You are a process analyst. The user is describing NEW steps or CHANGES to existing steps in a process flow. "
            "Focus ONLY on the new steps or modifications they are requesting. "
            "Do NOT describe the entire process—only what they want to add or change.\n\n"
            "Format your response as follows:\n"
            "**Planned changes**: [Brief summary of the NEW or CHANGED step(s), including who does it and in which system]\n\n"
            "**Clarifying Questions**: [If anything is unclear, ask 1-3 short clarifying questions as bullet points]\n\n"
            "Keep your response concise and focused on the delta. Use markdown formatting for emphasis."
        )

    msgs: List[Dict[str, Any]] = [{"role": "system", "content": system_content}]
    for m in messages_raw:
        msgs.append({"role": m.get("role", "user"), "content": m.get("content", "")})

    svc = get_chat_service()

    async def token_generator():
        async for chunk in svc.stream_chat(messages=msgs):
            yield chunk.encode("utf-8")

    return StreamingResponse(token_generator(), media_type="text/plain; charset=utf-8")


@app.post("/api/clarify-process/intent")
async def clarify_process_intent(payload: Dict[str, Any], user: Dict[str, Any] = Depends(verify_token)):
    return {"intent": "update_flow", "payload": payload}


@app.post("/api/process-taxonomy")
async def create_taxonomy(item: TaxonomyCreate, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    if item.level < 0 or item.level > 3:
        raise HTTPException(status_code=400, detail="Process level must be between L0 and L3")

    if item.parent_id:
        parent_row = await _get_process_row(conn, item.parent_id)
        if not parent_row:
            raise HTTPException(status_code=404, detail="Parent process not found")

        parent_level = int(parent_row.get("level", 0))
        expected_level = parent_level + 1
        if expected_level > 3:
            raise HTTPException(status_code=400, detail="Cannot create process beyond level L3")
        if item.level != expected_level:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid level for child process. Expected L{expected_level} under selected parent.",
            )

        parent_ctx = await _get_requester_role_context(conn, item.parent_id, user_id)
        if not parent_ctx["can_manage"]:
            await _raise_level_access_denied(conn, item.parent_id, "Creating a sub-process")

    item_id = str(uuid.uuid4())
    sort_order = item.sort_order
    if sort_order is None:
        sort_order = await _next_sibling_sort_order(conn, user_id, item.parent_id)
    await conn.execute(
        "INSERT INTO process_taxonomy (id, user_id, name, description, code, level, parent_id, sort_order) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (item_id, user_id, item.name, item.description, item.code, item.level, item.parent_id, sort_order),
    )
    await _ensure_creator_owner_assignment(conn, item_id)
    await conn.commit()
    payload = item.model_dump()
    payload["sort_order"] = sort_order
    payload["created_by"] = user_id
    return {"id": item_id, **payload}


@app.get("/api/process-taxonomy")
async def list_taxonomy(user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute("SELECT email FROM users WHERE id = ?", (user_id,)) as cursor:
        user_row = await cursor.fetchone()
    user_email = user_row["email"] if user_row else None

    async with conn.execute(
        """
        WITH RECURSIVE
        owner_seed(id) AS (
            SELECT id
            FROM process_taxonomy
            WHERE user_id = ?
            UNION
            SELECT process_id
            FROM process_assignments
            WHERE is_active = 1
              AND role = 'owner'
              AND (
                  user_id = ?
                  OR (? IS NOT NULL AND lower(user_email) = lower(?))
              )
        ),
        owner_scope(id) AS (
            SELECT id FROM owner_seed
            UNION
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN owner_scope os ON pt.parent_id = os.id
        )
        SELECT DISTINCT id FROM owner_scope
        """,
        (user_id, user_id, user_email, user_email),
    ) as cursor:
        owner_rows = await cursor.fetchall()
    owner_ids = {row["id"] for row in owner_rows}

    async with conn.execute(
        """
        WITH RECURSIVE
        delegator_seed(id) AS (
            SELECT process_id
            FROM process_assignments
            WHERE is_active = 1
              AND role = 'delegator'
              AND (
                  user_id = ?
                  OR (? IS NOT NULL AND lower(user_email) = lower(?))
              )
        ),
        delegator_scope(id) AS (
            SELECT id FROM delegator_seed
            UNION
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN delegator_scope ds ON pt.parent_id = ds.id
        )
        SELECT DISTINCT id FROM delegator_scope
        """,
        (user_id, user_email, user_email),
    ) as cursor:
        delegator_rows = await cursor.fetchall()
    delegator_ids = {row["id"] for row in delegator_rows}
    manageable_ids = owner_ids.union(delegator_ids)

    async with conn.execute(
        """
        WITH RECURSIVE
        seed(id) AS (
            SELECT id
            FROM process_taxonomy
            WHERE user_id = ?
            UNION
            SELECT process_id
            FROM process_assignments
            WHERE is_active = 1
              AND (
                  user_id = ?
                  OR (? IS NOT NULL AND lower(user_email) = lower(?))
              )
        ),
        -- Include descendants so assignments on parent processes flow down.
        down(id) AS (
            SELECT id FROM seed
            UNION
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN down d ON pt.parent_id = d.id
        ),
        -- Include ancestors to preserve full tree context in the UI.
        up(id, parent_id) AS (
            SELECT pt.id, pt.parent_id
            FROM process_taxonomy pt
            WHERE pt.id IN (SELECT id FROM down)
            UNION
            SELECT pt.id, pt.parent_id
            FROM process_taxonomy pt
            JOIN up ON up.parent_id = pt.id
        )
        SELECT DISTINCT
            pt.id,
            pt.user_id AS created_by,
            pt.name,
            pt.description,
            pt.code,
            pt.level,
            pt.parent_id,
            pt.sort_order
        FROM process_taxonomy pt
        WHERE pt.id IN (SELECT id FROM up)
        ORDER BY CASE WHEN pt.parent_id IS NULL THEN 0 ELSE 1 END, pt.parent_id, pt.sort_order, pt.created_at
        """,
        (user_id, user_id, user_email, user_email),
    ) as cursor:
        rows = await cursor.fetchall()
    by_parent: Dict[Optional[str], List[Dict[str, Any]]] = {}
    for row in rows:
        node = dict(row)
        by_parent.setdefault(node["parent_id"], []).append(node)

    for siblings in by_parent.values():
        siblings.sort(key=lambda node: (node.get("sort_order", 0), node.get("name", "")))

    child_counts = {parent_id: len(children) for parent_id, children in by_parent.items()}
    for siblings in by_parent.values():
        for node in siblings:
            node_id = node["id"]
            parent_id = node.get("parent_id")
            is_owner = node_id in owner_ids
            is_delegator = node_id in delegator_ids
            can_manage = node_id in manageable_ids
            is_leaf_delete = int(node.get("level", 0)) == 3 or child_counts.get(node_id, 0) == 0
            if is_owner:
                can_delete = True
            elif is_leaf_delete:
                can_delete = is_delegator
            else:
                can_delete = bool(parent_id) and parent_id in delegator_ids
            node["is_owner"] = is_owner
            node["is_delegator"] = is_delegator
            node["can_manage"] = can_manage
            node["can_delete"] = can_delete

    def build_tree(parent_id: Optional[str]) -> List[Dict[str, Any]]:
        children = by_parent.get(parent_id, [])
        for child in children:
            child["children"] = build_tree(child["id"])
        return children

    return build_tree(None)


@app.put("/api/process-taxonomy/{process_id}")
async def update_taxonomy(process_id: str, item: TaxonomyUpdate, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    await conn.execute(
        """
        UPDATE process_taxonomy
        SET name = COALESCE(?, name),
            description = COALESCE(?, description),
            code = COALESCE(?, code),
            sort_order = COALESCE(?, sort_order)
        WHERE id = ? AND user_id = ?
        """,
        (item.name, item.description, item.code, item.sort_order, process_id, user_id),
    )
    await conn.commit()
    return {"id": process_id, **item.model_dump()}


@app.delete("/api/process-taxonomy/{process_id}")
async def delete_taxonomy(process_id: str, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT id, user_id, parent_id, level FROM process_taxonomy WHERE id = ?",
        (process_id,),
    ) as cursor:
        root_row = await cursor.fetchone()
    if not root_row:
        raise HTTPException(status_code=404, detail="Process not found")

    requester_ctx = await _get_requester_role_context(conn, process_id, user_id)
    can_view_target = requester_ctx["has_access"] or await _can_view_process_in_tree(conn, process_id, user_id)
    if not can_view_target:
        raise HTTPException(status_code=403, detail="You do not have permission to delete this process")

    async with conn.execute(
        "SELECT COUNT(*) AS child_count FROM process_taxonomy WHERE parent_id = ?",
        (process_id,),
    ) as cursor:
        child_row = await cursor.fetchone()
    child_count = child_row["child_count"] if child_row else 0
    is_leaf_delete = int(root_row["level"]) == 3 or child_count == 0
    parent_id = root_row["parent_id"]

    if requester_ctx["is_owner"]:
        pass
    elif requester_ctx["is_delegator"]:
        if is_leaf_delete:
            pass
        elif parent_id:
            parent_ctx = await _get_requester_role_context(conn, parent_id, user_id)
            if not (parent_ctx["is_owner"] or parent_ctx["is_delegator"]):
                await _raise_level_access_denied(conn, parent_id, "Deleting this process")
        else:
            await _raise_level_access_denied(conn, process_id, "Deleting this process")
    else:
        reference_id = parent_id if is_leaf_delete and parent_id else process_id
        await _raise_level_access_denied(conn, reference_id, "Deleting this process")

    # Delete full subtree + dependent records to satisfy FK constraints.
    async with conn.execute(
        """
        WITH RECURSIVE subtree(id) AS (
            SELECT id FROM process_taxonomy WHERE id = ?
            UNION ALL
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN subtree s ON pt.parent_id = s.id
        )
        SELECT id FROM subtree
        """,
        (process_id,),
    ) as cursor:
        subtree_rows = await cursor.fetchall()
    subtree_ids = [row["id"] for row in subtree_rows]
    if not subtree_ids:
        raise HTTPException(status_code=404, detail="Process not found")

    subtree_placeholders = ",".join("?" for _ in subtree_ids)

    await conn.execute(
        f"DELETE FROM process_assignments WHERE process_id IN ({subtree_placeholders})",
        tuple(subtree_ids),
    )

    async with conn.execute(
        f"SELECT id FROM process_flows WHERE process_id IN ({subtree_placeholders})",
        tuple(subtree_ids),
    ) as cursor:
        flow_rows = await cursor.fetchall()
    flow_ids = [row["id"] for row in flow_rows]
    if flow_ids:
        flow_placeholders = ",".join("?" for _ in flow_ids)
        await conn.execute(
            f"DELETE FROM process_flow_versions WHERE flow_id IN ({flow_placeholders})",
            tuple(flow_ids),
        )

    await conn.execute(
        f"DELETE FROM process_flows WHERE process_id IN ({subtree_placeholders})",
        tuple(subtree_ids),
    )
    await conn.execute(
        f"DELETE FROM process_taxonomy WHERE id IN ({subtree_placeholders})",
        tuple(subtree_ids),
    )
    await conn.commit()
    return {"ok": True}


@app.post("/api/process-taxonomy/{process_id}/move")
async def move_taxonomy(process_id: str, payload: TaxonomyMoveRequest, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT id, level, parent_id, user_id FROM process_taxonomy WHERE id = ?",
        (process_id,),
    ) as cursor:
        source_row = await cursor.fetchone()
    if not source_row:
        raise HTTPException(status_code=404, detail="Process not found")

    source_parent_id = source_row["parent_id"]
    if source_row["user_id"] != user_id:
        if source_parent_id:
            source_parent_ctx = await _get_requester_role_context(conn, source_parent_id, user_id)
            if not source_parent_ctx["can_manage"]:
                await _raise_level_access_denied(conn, source_parent_id, "Moving this process")
        else:
            await _raise_level_access_denied(conn, process_id, "Moving this process")

    new_parent_id = payload.parent_id
    new_level = 0
    if new_parent_id is not None:
        async with conn.execute(
            "SELECT id, level FROM process_taxonomy WHERE id = ?",
            (new_parent_id,),
        ) as cursor:
            parent_row = await cursor.fetchone()
        if not parent_row:
            raise HTTPException(status_code=404, detail="New parent not found")
        target_parent_ctx = await _get_requester_role_context(conn, new_parent_id, user_id)
        if not target_parent_ctx["can_manage"]:
            await _raise_level_access_denied(conn, new_parent_id, "Moving this process")
        new_level = parent_row["level"] + 1

    if new_level > 3:
        raise HTTPException(status_code=400, detail="Cannot move process beyond level L3")

    subtree_height = await _subtree_height(conn, user_id, process_id)
    if new_level + subtree_height > 3:
        raise HTTPException(status_code=400, detail="Move would push descendants beyond level L3")

    sort_order = payload.sort_order
    if sort_order is None:
        sort_order = await _next_sibling_sort_order(conn, user_id, new_parent_id, exclude_id=process_id)

    await conn.execute(
        "UPDATE process_taxonomy SET parent_id = ?, sort_order = ? WHERE id = ? AND user_id = ?",
        (new_parent_id, sort_order, process_id, user_id),
    )
    await _update_subtree_levels(conn, user_id, process_id, new_level)
    await conn.commit()
    return {"id": process_id, "parent_id": new_parent_id, "sort_order": sort_order, "level": new_level}


@app.post("/api/process-taxonomy/reorder")
async def reorder_taxonomy(payload: TaxonomyReorderRequest, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    for item in payload.items:
        async with conn.execute(
            "SELECT 1 FROM process_taxonomy WHERE id = ? AND user_id = ?",
            (item.id, user_id),
        ) as cursor:
            exists = await cursor.fetchone()
        if not exists:
            raise HTTPException(status_code=404, detail=f"Process not found: {item.id}")

        await conn.execute(
            "UPDATE process_taxonomy SET sort_order = ? WHERE id = ? AND user_id = ?",
            (item.sort_order, item.id, user_id),
        )
    await conn.commit()
    return {"ok": True, "updated": len(payload.items)}


@app.post("/api/process-taxonomy/{process_id}/duplicate")
async def duplicate_taxonomy(process_id: str, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT name, description, code, level, parent_id FROM process_taxonomy WHERE id = ? AND user_id = ?",
        (process_id, user_id),
    ) as cursor:
        row = await cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Process not found")
    sort_order = await _next_sibling_sort_order(conn, user_id, row["parent_id"])
    new_id = str(uuid.uuid4())
    await conn.execute(
        "INSERT INTO process_taxonomy (id, user_id, name, description, code, level, parent_id, sort_order) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (new_id, user_id, f"{row['name']} (Copy)", row["description"], row["code"], row["level"], row["parent_id"], sort_order),
    )
    await _ensure_creator_owner_assignment(conn, new_id)
    await conn.commit()
    return {"id": new_id, "name": f"{row['name']} (Copy)", "sort_order": sort_order}


@app.post("/api/process-assignments")
async def create_process_assignment(
    payload: ProcessAssignmentCreateRequest,
    user_id: str = Depends(get_current_user_id),
):
    conn = db.get_connection()
    process_row = await _get_process_row(conn, payload.process_id)
    if not process_row:
        raise HTTPException(status_code=404, detail="Process not found")

    role = _normalize_role(payload.role)
    requester_ctx = await _get_requester_role_context(conn, payload.process_id, user_id)
    if not requester_ctx["can_manage"]:
        raise HTTPException(status_code=403, detail="Only owners or delegators can assign users")
    if role == "owner" and not requester_ctx["is_owner"]:
        raise HTTPException(status_code=403, detail="Only owners can assign owner role")

    target_email = payload.user_email.lower()
    async with conn.execute(
        "SELECT id FROM users WHERE lower(email) = lower(?) LIMIT 1",
        (target_email,),
    ) as cursor:
        target_user = await cursor.fetchone()
    target_user_id = target_user["id"] if target_user else None

    await conn.execute(
        """
        UPDATE process_assignments
        SET is_active = 0
        WHERE process_id = ? AND lower(user_email) = lower(?) AND is_active = 1
        """,
        (payload.process_id, target_email),
    )
    assignment_id = str(uuid.uuid4())
    await conn.execute(
        """
        INSERT INTO process_assignments (id, process_id, user_email, user_id, role, assigned_by, is_active)
        VALUES (?, ?, ?, ?, ?, ?, 1)
        """,
        (assignment_id, payload.process_id, target_email, target_user_id, role, user_id),
    )
    await conn.commit()

    return {
        "id": assignment_id,
        "process_id": payload.process_id,
        "user_email": target_email,
        "user_id": target_user_id,
        "role": role,
        "assigned_by": user_id,
        "is_active": True,
    }


@app.get("/api/process-assignments/process/{process_id}")
async def list_process_assignments(process_id: str, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    process_row = await _get_process_row(conn, process_id)
    if not process_row:
        raise HTTPException(status_code=404, detail="Process not found")

    requester_ctx = await _get_requester_role_context(conn, process_id, user_id)
    if not requester_ctx["has_access"] and not await _can_view_process_in_tree(conn, process_id, user_id):
        raise HTTPException(status_code=403, detail="You do not have permission to view assignments for this process")

    assignments = await _get_effective_assignments(conn, process_id)
    return assignments


@app.delete("/api/process-assignments/{assignment_id}")
async def delete_process_assignment(assignment_id: str, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        """
        SELECT id, process_id, role
        FROM process_assignments
        WHERE id = ? AND is_active = 1
        """,
        (assignment_id,),
    ) as cursor:
        assignment_row = await cursor.fetchone()
    if not assignment_row:
        raise HTTPException(status_code=404, detail="Assignment not found")

    process_id = assignment_row["process_id"]
    requester_ctx = await _get_requester_role_context(conn, process_id, user_id)
    if not requester_ctx["can_manage"]:
        raise HTTPException(status_code=403, detail="Only owners or delegators can remove assignments")

    if assignment_row["role"] == "owner":
        if not requester_ctx["is_owner"]:
            raise HTTPException(status_code=403, detail="Only owners can remove owner assignments")
        async with conn.execute(
            """
            SELECT COUNT(*) AS owner_count
            FROM process_assignments
            WHERE process_id = ? AND role = 'owner' AND is_active = 1
            """,
            (process_id,),
        ) as cursor:
            owner_row = await cursor.fetchone()
        owner_count = owner_row["owner_count"] if owner_row else 0
        if owner_count <= 1:
            raise HTTPException(status_code=400, detail="Cannot remove the only owner from this process")

    await conn.execute("UPDATE process_assignments SET is_active = 0 WHERE id = ?", (assignment_id,))
    await conn.commit()
    return {"ok": True}


@app.post("/api/process-flows")
async def create_flow(item: FlowCreate, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    flow_id = str(uuid.uuid4())
    flow_json = json.dumps(item.flow_data)
    await conn.execute(
        "INSERT INTO process_flows (id, user_id, process_id, title, description, flow_data) VALUES (?, ?, ?, ?, ?, ?)",
        (flow_id, user_id, item.process_id, item.title, item.description, flow_json),
    )
    await conn.execute(
        "INSERT INTO process_flow_versions (id, flow_id, version, flow_data) VALUES (?, ?, ?, ?)",
        (str(uuid.uuid4()), flow_id, 1, flow_json),
    )
    await conn.commit()
    return {"id": flow_id, **item.model_dump()}


@app.get("/api/process-flows/process/{process_id}")
async def list_flows_by_process(process_id: str, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT id, process_id, title, description, flow_data, version FROM process_flows WHERE user_id = ? AND process_id = ?",
        (user_id, process_id),
    ) as cursor:
        rows = await cursor.fetchall()
    out = []
    for row in rows:
        item = dict(row)
        item["flow_data"] = json.loads(item["flow_data"])
        out.append(item)
    return out


@app.put("/api/process-flows/{flow_id}")
async def update_flow(flow_id: str, item: FlowUpdate, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    flow_json = json.dumps(item.flow_data)
    await conn.execute(
        """
        UPDATE process_flows
        SET title = COALESCE(?, title),
            description = COALESCE(?, description),
            flow_data = ?,
            version = version + 1,
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ? AND user_id = ?
        """,
        (item.title, item.description, flow_json, flow_id, user_id),
    )
    async with conn.execute("SELECT version FROM process_flows WHERE id = ? AND user_id = ?", (flow_id, user_id)) as cursor:
        row = await cursor.fetchone()
    version = row["version"] if row else 1
    await conn.execute(
        "INSERT INTO process_flow_versions (id, flow_id, version, flow_data) VALUES (?, ?, ?, ?)",
        (str(uuid.uuid4()), flow_id, version, flow_json),
    )
    await conn.commit()
    return {"id": flow_id, "version": version, **item.model_dump()}


@app.get("/api/process-flows/{flow_id}/versions")
async def flow_versions(flow_id: str, _user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT version, created_at FROM process_flow_versions WHERE flow_id = ? ORDER BY version DESC",
        (flow_id,),
    ) as cursor:
        rows = await cursor.fetchall()
    return [{"version_number": r["version"], "created_at": r["created_at"]} for r in rows]


@app.get("/api/process-flows/{flow_id}/versions/{version}/data")
async def flow_version_data(flow_id: str, version: int, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT flow_data FROM process_flow_versions WHERE flow_id = ? AND version = ?",
        (flow_id, version),
    ) as cursor:
        row = await cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Version not found")
    return json.loads(row["flow_data"])


@app.post("/api/process-flows/{flow_id}/versions/{version}/restore")
async def restore_version(flow_id: str, version: int, user_id: str = Depends(get_current_user_id)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT flow_data FROM process_flow_versions WHERE flow_id = ? AND version = ?",
        (flow_id, version),
    ) as cursor:
        row = await cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Version not found")
    flow_data = json.loads(row["flow_data"])

    # Save as a new version
    async with conn.execute(
        "SELECT COALESCE(MAX(version), 0) FROM process_flow_versions WHERE flow_id = ?",
        (flow_id,),
    ) as cursor:
        max_row = await cursor.fetchone()
    new_version = (max_row[0] if max_row else 0) + 1
    flow_json = json.dumps(flow_data)
    await conn.execute(
        "INSERT INTO process_flow_versions (id, flow_id, version, flow_data) VALUES (?, ?, ?, ?)",
        (str(uuid.uuid4()), flow_id, new_version, flow_json),
    )
    # Also update the main flow record
    await conn.execute(
        "UPDATE process_flows SET flow_data = ?, updated_at = ? WHERE id = ?",
        (flow_json, datetime.utcnow().isoformat(), flow_id),
    )
    await conn.commit()
    return {"version": new_version, "flow_data": flow_data}


@app.post("/api/flow/rollback-patch")
async def rollback_patch(payload: Dict[str, Any], user_id: str = Depends(get_current_user_id)):
    return {"ok": True, "payload": payload}


@app.post("/api/export/excel")
async def export_excel(payload: Dict[str, Any], user_id: str = Depends(get_current_user_id)):
    import io
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    selected_l0_ids: List[str] = payload.get("selected_l0_ids", [])
    taxonomy_attrs: List[str] = payload.get("taxonomy_attributes", [])
    flow_attrs: List[str] = payload.get("flow_attributes", [])
    include_flow_nodes: bool = payload.get("include_flow_nodes", True)
    node_attrs: List[str] = payload.get("node_attributes", [])

    if not selected_l0_ids:
        raise HTTPException(status_code=400, detail="No L0 processes selected")

    conn = db.get_connection()

    # Collect all taxonomy rows for the user, then filter to selected L0 subtrees
    async with conn.execute(
        "SELECT id, parent_id, name, description, level, code, sort_order, created_at "
        "FROM process_taxonomy WHERE user_id = ? ORDER BY level, sort_order",
        (user_id,),
    ) as cursor:
        all_rows = [dict(r) for r in await cursor.fetchall()]

    by_id = {r["id"]: r for r in all_rows}
    children_of: Dict[Optional[str], List[dict]] = {}
    for r in all_rows:
        children_of.setdefault(r["parent_id"], []).append(r)

    # Walk subtrees of selected L0s
    def collect_subtree(root_id: str) -> List[dict]:
        items = []
        node = by_id.get(root_id)
        if not node:
            return items
        items.append(node)
        for child in children_of.get(root_id, []):
            items.extend(collect_subtree(child["id"]))
        return items

    included: List[dict] = []
    for l0_id in selected_l0_ids:
        included.extend(collect_subtree(l0_id))

    # Build hierarchy lookup (walk parent chain for each item)
    def get_hierarchy(item_id: str) -> Dict[str, str]:
        h = {}
        cur = item_id
        while cur:
            node = by_id.get(cur)
            if not node:
                break
            h[node["level"]] = node["name"]
            cur = node["parent_id"]
        return {"l0": h.get(0, ""), "l1": h.get(1, ""), "l2": h.get(2, "")}

    # Fetch user info for owner display
    async with conn.execute("SELECT id, email, username FROM users WHERE id = ?", (user_id,)) as cursor:
        user_row = await cursor.fetchone()
    owner_display = (dict(user_row)["username"] if user_row else "N/A")

    # Fetch all flows for included processes
    process_ids = [r["id"] for r in included]
    flows_by_process: Dict[str, dict] = {}
    if process_ids:
        placeholders = ",".join("?" for _ in process_ids)
        async with conn.execute(
            f"SELECT id, process_id, user_id, title, description, flow_data, version, "
            f"created_at, updated_at FROM process_flows WHERE process_id IN ({placeholders})",
            process_ids,
        ) as cursor:
            for row in await cursor.fetchall():
                flows_by_process[row["process_id"]] = dict(row)

    # --- Build workbook ---
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ---- Sheet 1: Taxonomy ----
    ws_tax = wb.active
    ws_tax.title = "Taxonomy"

    tax_col_map = [
        ("id", "ID"),
        ("parent_id", "Parent ID"),
        ("l0", "L0 Name"),
        ("l1", "L1 Name"),
        ("l2", "L2 Name"),
        ("level", "Level"),
        ("name", "Process Name"),
        ("description", "Description"),
        ("updated_at", "Updated At"),
        ("role", "Owner(s)"),
    ]
    tax_cols = [(key, label) for key, label in tax_col_map if key in taxonomy_attrs]

    for ci, (_, label) in enumerate(tax_cols, 1):
        cell = ws_tax.cell(row=1, column=ci, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    def get_taxonomy_value(item: dict, flow: Optional[dict], hierarchy: Dict[str, str], key: str):
        if key == "l0":
            return hierarchy.get("l0", "")
        if key == "l1":
            return hierarchy.get("l1", "")
        if key == "l2":
            return hierarchy.get("l2", "")
        if key == "role":
            return owner_display
        if key == "updated_at":
            return (flow["updated_at"] if flow else item.get("created_at")) or ""
        return item.get(key, "")

    for ri, item in enumerate(included, 2):
        hierarchy = get_hierarchy(item["id"])
        flow = flows_by_process.get(item["id"])
        for ci, (key, _) in enumerate(tax_cols, 1):
            val = get_taxonomy_value(item, flow, hierarchy, key)
            cell = ws_tax.cell(row=ri, column=ci, value=val if val is not None else "")
            cell.border = thin_border

    for ci in range(1, len(tax_cols) + 1):
        ws_tax.column_dimensions[ws_tax.cell(row=1, column=ci).column_letter].width = 20

    # ---- Sheet 2: Flow Nodes (if requested) ----
    if include_flow_nodes:
        ws_nodes = wb.create_sheet("Flow Nodes")

        flow_col_map = [
            ("process_name", "Process Name"),
            ("title", "Flow Title"),
            ("description", "Flow Description"),
            ("created_by", "Created By"),
            ("status", "Status"),
            ("version", "Flow Version"),
            ("created_at", "Flow Created At"),
            ("updated_at", "Flow Updated At"),
        ]
        flow_cols = [(key, label) for key, label in flow_col_map if key in flow_attrs]

        node_col_map = [
            ("node_type", "Node Type"),
            ("node_label", "Process Step"),
            ("node_owner", "Owner"),
            ("node_system", "Tool/System"),
            ("node_automation", "Manual/Automated"),
        ]
        active_node_keys = set(node_attrs)
        node_cols = [(key, label) for key, label in node_col_map if key in active_node_keys]
        expanded_cols = tax_cols + flow_cols + node_cols
        taxonomy_keys = {k for k, _ in tax_cols}

        for ci, (_, label) in enumerate(expanded_cols, 1):
            cell = ws_nodes.cell(row=1, column=ci, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        row_idx = 2
        for item in included:
            flow = flows_by_process.get(item["id"])
            if not flow:
                continue
            raw = flow.get("flow_data") or "{}"
            try:
                fd = json.loads(raw) if isinstance(raw, str) else raw
            except json.JSONDecodeError:
                continue
            nodes_list = fd.get("nodes", [])
            if not nodes_list:
                continue

            for node in nodes_list:
                nd = node.get("data", node)
                hierarchy = get_hierarchy(item["id"])
                flow_vals = {
                    "process_name": item["name"],
                    "title": flow.get("title", ""),
                    "description": flow.get("description", ""),
                    "created_by": flow.get("user_id", ""),
                    "status": flow.get("status", ""),
                    "version": flow.get("version", ""),
                    "created_at": flow.get("created_at", ""),
                    "updated_at": flow.get("updated_at", ""),
                }
                node_vals = {
                    "node_type": node.get("type") or nd.get("type", "process"),
                    "node_label": nd.get("label", ""),
                    "node_owner": nd.get("owner", ""),
                    "node_system": nd.get("system", ""),
                    "node_automation": nd.get("manualOrAutomated", ""),
                }
                for ci, (key, _) in enumerate(expanded_cols, 1):
                    if key in taxonomy_keys:
                        value = get_taxonomy_value(item, flow, hierarchy, key)
                    elif key in flow_vals:
                        value = flow_vals.get(key, "")
                    else:
                        value = node_vals.get(key, "")
                    cell = ws_nodes.cell(row=row_idx, column=ci, value=value)
                    cell.border = thin_border
                row_idx += 1

        for ci in range(1, len(expanded_cols) + 1):
            ws_nodes.column_dimensions[ws_nodes.cell(row=1, column=ci).column_letter].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="process_export.xlsx"'},
    )


@app.get("/api/dashboard/{user_id}")
async def dashboard(user_id: str, current_user: Dict[str, Any] = Depends(verify_token)):
    def _flow_has_start_and_end_nodes(flow_data_raw: Any) -> bool:
        """A process is completed only if its flow contains both start and end nodes."""
        if flow_data_raw is None:
            return False

        flow_data = flow_data_raw
        if isinstance(flow_data_raw, str):
            try:
                flow_data = json.loads(flow_data_raw)
            except Exception:
                return False

        if not isinstance(flow_data, dict):
            return False

        nodes = flow_data.get("nodes")
        if not isinstance(nodes, list) or not nodes:
            return False

        has_start = False
        has_end = False
        for node in nodes:
            if not isinstance(node, dict):
                continue
            node_type = str(node.get("type") or node.get("data", {}).get("type") or "").lower()
            if node_type == "start":
                has_start = True
            elif node_type == "end":
                has_end = True
            if has_start and has_end:
                return True

        return False

    conn = db.get_connection()
    async with conn.execute("SELECT email FROM users WHERE id = ?", (user_id,)) as cursor:
        user_row = await cursor.fetchone()
    user_email = (user_row["email"] if user_row else None)

    async with conn.execute(
        """
        WITH RECURSIVE
        seed(id) AS (
            SELECT id
            FROM process_taxonomy
            WHERE user_id = ?
            UNION
            SELECT process_id
            FROM process_assignments
            WHERE is_active = 1
              AND (
                  user_id = ?
                  OR (? IS NOT NULL AND lower(user_email) = lower(?))
              )
        ),
        accessible(id) AS (
            SELECT id FROM seed
            UNION
            SELECT pt.id
            FROM process_taxonomy pt
            JOIN accessible a ON pt.parent_id = a.id
        )
        SELECT DISTINCT pt.id
        FROM process_taxonomy pt
        WHERE pt.id IN (SELECT id FROM accessible)
          AND pt.level = 3
        """,
        (user_id, user_id, user_email, user_email),
    ) as cursor:
        process_rows = await cursor.fetchall()
    process_ids = [row["id"] for row in process_rows]

    completed_ids = []
    for pid in process_ids:
        async with conn.execute(
            """
            SELECT flow_data FROM process_flows
            WHERE process_id = ?
            ORDER BY updated_at DESC, created_at DESC
            LIMIT 1
            """,
            (pid,),
        ) as cursor:
            row = await cursor.fetchone()
            if row and _flow_has_start_and_end_nodes(row["flow_data"]):
                completed_ids.append(pid)

    total = len(process_ids)
    pct = (len(completed_ids) / total * 100) if total > 0 else 0
    return {
        "stats": {"total_processes": total, "completion_percentage": pct},
        "completed_process_ids": completed_ids,
    }


@app.get("/api/processes/{process_id}/sop-status")
async def sop_status(process_id: str, user: Dict[str, Any] = Depends(verify_token)):
    conn = db.get_connection()
    async with conn.execute(
        "SELECT sop_status, sop_draft_url, sop_final_url, sop_final_pdf_url, "
        "sop_generated_at, sop_finalized_at FROM process_taxonomy WHERE id = ?",
        (process_id,),
    ) as cursor:
        row = await cursor.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Process not found")
    result = {
        "status": row["sop_status"],
        "draft_url": row["sop_draft_url"],
        "final_url": row["sop_final_url"],
        "final_pdf_url": row["sop_final_pdf_url"],
        "has_recording_data": False,
        "generated_at": row["sop_generated_at"],
        "finalized_at": row["sop_finalized_at"],
        "flow_out_of_sync": False,
        "flow_last_synced_at": None,
    }
    storage_service.sign_sop_urls(result)
    return result


async def get_process_hierarchy(conn, process_id: str) -> Dict[str, str]:
    """Walk up the taxonomy tree to get L0-L3 hierarchy names."""
    hierarchy = {"l0_name": "", "l1_name": "", "l2_name": "", "l3_name": ""}
    current_id = process_id
    while current_id:
        async with conn.execute(
            "SELECT id, name, level, parent_id FROM process_taxonomy WHERE id = ?",
            (current_id,),
        ) as cursor:
            row = await cursor.fetchone()
        if not row:
            break
        hierarchy[f"l{row['level']}_name"] = row["name"]
        current_id = row["parent_id"]
    return hierarchy


def _sanitize_filename(name: str) -> str:
    """Remove characters that are invalid in filenames."""
    import re as _re
    return _re.sub(r'[\\/:*?"<>|]', "_", name).strip()[:100]


class SingleSOPExportRequest(BaseModel):
    include_ai_suggestions: bool = True
    flow_screenshot: Optional[str] = None


@app.post("/api/export/sop/{process_id}")
async def export_sop(
    process_id: str,
    export_request: Optional[SingleSOPExportRequest] = None,
    user_id: str = Depends(get_current_user_id),
):
    """Generate a SOP Word document and return it as a direct download."""
    from fastapi.responses import Response

    include_ai = export_request.include_ai_suggestions if export_request else True
    flow_screenshot = export_request.flow_screenshot if export_request else None

    conn = db.get_connection()

    async with conn.execute(
        "SELECT id, user_id, parent_id, level, name, description, code, created_at "
        "FROM process_taxonomy WHERE id = ?",
        (process_id,),
    ) as cursor:
        process_row = await cursor.fetchone()
    if not process_row:
        raise HTTPException(status_code=404, detail="Process not found")
    process = dict(process_row)

    async with conn.execute(
        "SELECT id, process_id, user_id AS created_by, title, description, flow_data, "
        "version, created_at, updated_at FROM process_flows "
        "WHERE process_id = ? ORDER BY updated_at DESC LIMIT 1",
        (process_id,),
    ) as cursor:
        flow_row = await cursor.fetchone()

    flow = None
    flow_data = None
    if flow_row:
        flow = dict(flow_row)
        raw = flow.get("flow_data") or "{}"
        if isinstance(raw, str):
            try:
                flow_data = json.loads(raw)
            except json.JSONDecodeError:
                flow_data = {}
        else:
            flow_data = raw
        flow["flow_data"] = flow_data

    hierarchy = await get_process_hierarchy(conn, process_id)

    # Look up the process creator as the owner
    async with conn.execute(
        "SELECT email, username FROM users WHERE id = ?",
        (process.get("user_id"),),
    ) as cursor:
        user_row = await cursor.fetchone()
    assignments: List[Dict[str, Any]] = []
    if user_row:
        assignments.append({
            "role": "owner",
            "full_name": user_row["username"],
            "user_email": user_row["email"],
        })

    # Inject updated_at and version into the process dict for the title page
    if flow:
        process["updated_at"] = flow.get("updated_at") or flow.get("created_at")
        process["version"] = flow.get("version", 1)

    ai_suggestions = None
    if include_ai:
        try:
            sop_ai = get_sop_ai_service()
            result = await sop_ai.generate_sop_suggestions(
                process=process,
                flow_data=flow_data,
                hierarchy=hierarchy,
                assignments=assignments,
            )
            ai_suggestions = [
                {
                    "suggestion_type": s.suggestion_type.value,
                    "section": s.section,
                    "original_text": s.original_text,
                    "suggested_text": s.suggested_text,
                    "comment_text": s.comment_text,
                    "node_id": s.node_id,
                    "field_name": s.field_name,
                }
                for s in result.suggestions
            ]
            logger.info("Generated %d AI SOP suggestions for %s", len(ai_suggestions), process_id)
        except Exception as ai_err:
            logger.warning("AI suggestions failed for SOP %s: %s", process_id, ai_err)

    from services.sop_generator import SOPGenerator
    generator = SOPGenerator()
    doc_buffer = await generator.generate_sop(
        process=process,
        flow=flow,
        assignments=assignments,
        hierarchy=hierarchy,
        ai_suggestions=ai_suggestions,
        flow_screenshot=flow_screenshot,
        storage_service=storage_service,
    )

    doc_bytes = doc_buffer.getvalue()
    safe_name = _sanitize_filename(process.get("name", "Process"))
    date_str = datetime.utcnow().strftime("%y%m%d")
    filename = f"{safe_name} - {date_str} vDraft.docx"

    return Response(
        content=doc_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


SCREENSHOT_DIR = os.path.join(os.path.dirname(__file__), "screenshots")
os.makedirs(SCREENSHOT_DIR, exist_ok=True)


@app.post("/api/process-flows/{flow_id}/upload-screenshot")
async def upload_screenshot(
    flow_id: str,
    file: UploadFile = File(...),
    node_id: Optional[str] = Form(None),
    user_id: str = Depends(get_current_user_id),
):
    import hashlib
    content = await file.read()
    ext = os.path.splitext(file.filename or "img.png")[1] or ".png"
    file_hash = hashlib.sha256(content).hexdigest()[:12]
    filename = f"{flow_id}_{file_hash}{ext}"
    filepath = os.path.join(SCREENSHOT_DIR, filename)
    with open(filepath, "wb") as f:
        f.write(content)
    url = f"/api/screenshots/{filename}"
    return {"key": url, "url": url, "node_id": node_id}


@app.get("/api/screenshots/{filename}")
async def serve_screenshot(filename: str):
    from fastapi.responses import FileResponse
    filepath = os.path.join(SCREENSHOT_DIR, filename)
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail="Screenshot not found")
    return FileResponse(filepath)


@app.get("/api/process-flows/{flow_id}/screenshots")
async def list_screenshots(flow_id: str, user_id: str = Depends(get_current_user_id)):
    prefix = f"{flow_id}_"
    shots = []
    if os.path.isdir(SCREENSHOT_DIR):
        for fname in os.listdir(SCREENSHOT_DIR):
            if fname.startswith(prefix):
                shots.append({"url": f"/api/screenshots/{fname}", "captured_at": None, "note": "Local upload"})
    return {"screenshots": shots}


@app.delete("/api/process-flows/{flow_id}/screenshots")
async def delete_screenshot(flow_id: str, payload: Dict[str, Any] = {}, user_id: str = Depends(get_current_user_id)):
    key = payload.get("key", "")
    filename = key.split("/")[-1] if "/" in key else key
    filepath = os.path.join(SCREENSHOT_DIR, filename)
    if os.path.exists(filepath):
        os.remove(filepath)
    return {"ok": True}


@app.get("/api/storage/{path:path}")
async def serve_local_storage(path: str):
    """Serve files from local_storage/ directory (SOP docs, etc.)."""
    from fastapi.responses import FileResponse
    full = storage_service.get_local_path(path)
    if not full.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(str(full))


@app.post("/api/contact")
async def contact(payload: Dict[str, Any]):
    return {"message": "Thanks, we received your message.", "payload": payload}


if __name__ == "__main__":
    uvicorn.run("app:app", host="0.0.0.0", port=int(os.getenv("PORT", "8000")), reload=True)
